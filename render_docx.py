import docx
from docx.shared import Inches
import requests
import os
import time
from dotenv import load_dotenv

# 1. Load the keys
load_dotenv()

PEXELS_HEADERS = {
    "Authorization": os.getenv("PEXELS_API_KEY"),
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}

WIKI_HEADERS = {
    "User-Agent": "GeminiDocAgent/3.0 (LearningProject; hybrid-engine)"
}

def get_wiki_image_url(search_term, used_urls):
    """Searches Wikipedia for factual, encyclopedic images."""
    print(f" -> Searching Wikipedia (Factual) for: '{search_term}'")
    url = "https://en.wikipedia.org/w/api.php"
    params = {
        "action": "query",
        "generator": "search",
        "gsrsearch": search_term,
        "gsrlimit": 1,
        "prop": "pageimages",
        "piprop": "thumbnail",
        "pithumbsize": 800,
        "format": "json"
    }
    
    try:
        response = requests.get(url, params=params, headers=WIKI_HEADERS, timeout=10)
        response.raise_for_status()
        data = response.json()
        
        pages = data.get("query", {}).get("pages", {})
        for page_id in pages:
            if "thumbnail" in pages[page_id]:
                img_url = pages[page_id]["thumbnail"]["source"]
                # Duplicate Check
                if img_url not in used_urls:
                    used_urls.add(img_url)
                    return img_url
    except Exception as e:
        print(f" -> Wiki API Error: {e}")
    return None

def get_pexels_image_url(search_term, used_urls):
    """Searches Pexels for high-quality stock photography."""
    if not os.getenv("PEXELS_API_KEY"):
        return None

    print(f" -> Searching Pexels (Stock) for: '{search_term}'")
    url = "https://api.pexels.com/v1/search"
    params = {"query": search_term, "per_page": 5} 
    
    time.sleep(2) # Be polite to Pexels
    
    try:
        response = requests.get(url, headers=PEXELS_HEADERS, params=params, timeout=10)
        response.raise_for_status() 
        data = response.json()
        
        if data.get('photos'):
            for photo in data['photos']:
                img_url = photo['src']['large']
                # Duplicate Check
                if img_url not in used_urls:
                    used_urls.add(img_url)
                    return img_url
    except Exception as e:
        print(f" -> Pexels API Error: {e}")
        
    return None

def create_word_doc(data, filename="output.docx"):
    """Takes JSON data and builds a formatted Word document."""
    print(f"Drafting Word document: {filename}...")
    
    doc = docx.Document()
    
    title = data.get("title", "AI Generated Report")
    doc.add_heading(title, 0)
    
    # 2. Create the memory bank for the Word doc
    used_images = set()
    
    for item in data.get("content", []):
        if item["type"] == "heading":
            doc.add_heading(item["text"], level=1)
            
        elif item["type"] == "paragraph":
            doc.add_paragraph(item["text"])
            
        elif item["type"] == "bullet_list":
            for point in item["items"]:
                doc.add_paragraph(point, style='List Bullet')
                
        elif item["type"] == "image":
            search_term = item.get("search_term")
            fallback_term = item.get("fallback_search_term")
            image_type = item.get("image_type", "stock") # Default to stock
            
            if search_term:
                img_url = None
                
                # THE HYBRID ROUTING LOGIC
                if image_type == "factual":
                    img_url = get_wiki_image_url(search_term, used_images)
                
                if not img_url:
                    img_url = get_pexels_image_url(search_term, used_images)
                
                if not img_url and fallback_term:
                    print(f" -> No exact matches. Trying fallback: '{fallback_term}'")
                    img_url = get_pexels_image_url(fallback_term, used_images)
                
                if img_url:
                    try:
                        # Download using the new requests logic
                        response = requests.get(img_url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=15, stream=True)
                        response.raise_for_status()
                        
                        with open("temp_doc_img.jpg", 'wb') as out_file:
                            out_file.write(response.content)
                            
                        # Insert into Word doc
                        doc.add_picture("temp_doc_img.jpg", width=Inches(5.0))
                        os.remove("temp_doc_img.jpg")
                        
                    except Exception as e:
                        print(f" -> Could not download image: {e}")
                        doc.add_paragraph(f"[Image Placeholder: {search_term}]")
                else:
                    print(f" -> No unique images found. Leaving text space.")
                    doc.add_paragraph(f"[Image intended: {search_term}]")
                    
    doc.save(filename)
    print("Word document complete!")