from pptx import Presentation
from pptx.util import Inches, Pt
import requests
import os
import time
from dotenv import load_dotenv

load_dotenv()

PEXELS_HEADERS = {
    "Authorization": os.getenv("PEXELS_API_KEY"),
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}

WIKI_HEADERS = {
    "User-Agent": "GeminiDocAgent/3.0 (LearningProject; hybrid-engine)"
}

def get_wiki_image_url(search_term, used_urls):
    """Searches Wikipedia for factual, encyclopedic images."""
    print(f" -> Searching Wikipedia (Factual) for: '{search_term}'")
    url = "https://en.wikipedia.org/w/api.php"
    params = {
        "action": "query",
        "generator": "search",
        "gsrsearch": search_term,
        "gsrlimit": 1,
        "prop": "pageimages",
        "piprop": "thumbnail",
        "pithumbsize": 800,
        "format": "json"
    }
    
    try:
        response = requests.get(url, params=params, headers=WIKI_HEADERS, timeout=10)
        response.raise_for_status()
        data = response.json()
        
        pages = data.get("query", {}).get("pages", {})
        for page_id in pages:
            if "thumbnail" in pages[page_id]:
                img_url = pages[page_id]["thumbnail"]["source"]
                if img_url not in used_urls:
                    used_urls.add(img_url)
                    return img_url
    except Exception as e:
        print(f" -> Wiki API Error: {e}")
    return None

def get_pexels_image_url(search_term, used_urls):
    """Searches Pexels for high-quality stock photography."""
    if not os.getenv("PEXELS_API_KEY"):
        return None

    print(f" -> Searching Pexels (Stock) for: '{search_term}'")
    url = "https://api.pexels.com/v1/search"
    params = {"query": search_term, "per_page": 5} 
    
    time.sleep(2) # Be polite to Pexels
    
    try:
        response = requests.get(url, headers=PEXELS_HEADERS, params=params, timeout=10)
        response.raise_for_status() 
        data = response.json()
        
        if data.get('photos'):
            for photo in data['photos']:
                img_url = photo['src']['large']
                if img_url not in used_urls:
                    used_urls.add(img_url)
                    return img_url
    except Exception as e:
        print(f" -> Pexels API Error: {e}")
        
    return None

def create_ppt_file(data, filename="output.pptx"):
    print(f"Drafting PowerPoint: {filename}...")
    prs = Presentation()

    slides_data = data.get("slides", [])
    used_images = set() 

    for slide_data in slides_data:
        if slide_data["type"] == "title_slide":
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = slide_data.get("title", "")
            slide.placeholders[1].text = slide_data.get("subtitle", "")

        elif slide_data["type"] == "content_slide":
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            
            slide.shapes.title.text = slide_data.get("title", "")

            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.word_wrap = True 
            
            search_term = slide_data.get("search_term")
            fallback_term = slide_data.get("fallback_search_term")
            image_type = slide_data.get("image_type", "stock") # Default to stock if missing
            
            if search_term:
                body_shape.left = Inches(0.5)
                body_shape.width = Inches(4.5)
                body_shape.top = Inches(1.8)
                body_shape.height = Inches(5.0)
                
                img_url = None
                
                # THE HYBRID ROUTING LOGIC
                if image_type == "factual":
                    # Try Wiki first for real people/places
                    img_url = get_wiki_image_url(search_term, used_images)
                
                if not img_url:
                    # If it wasn't factual, OR if Wikipedia didn't have it, try Pexels
                    img_url = get_pexels_image_url(search_term, used_images)
                
                if not img_url and fallback_term:
                    # The Ultimate Fallback
                    print(f" -> No exact matches. Trying fallback: '{fallback_term}'")
                    img_url = get_pexels_image_url(fallback_term, used_images)
                
                if img_url:
                    try:
                        response = requests.get(img_url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=15, stream=True)
                        response.raise_for_status() 

                        with open("temp_ppt_img.jpg", 'wb') as out_file:
                            out_file.write(response.content)
            
                        slide.shapes.add_picture("temp_ppt_img.jpg", left=Inches(5.2), top=Inches(1.8), height=Inches(4.5))
                        os.remove("temp_ppt_img.jpg")
                        
                    except Exception as e:
                        print(f" -> Could not download image: {e}")
                else:
                    print(f" -> No unique images found. Leaving text space.")
            else:
                body_shape.width = Inches(9.0)

            bullets = slide_data.get("bullets", [])
            for i, bullet in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = bullet
                p.font.size = Pt(20)
                p.space_after = Pt(14) 

    prs.save(filename)
    print("PowerPoint complete!")