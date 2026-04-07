import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_file(data, filename="output.xlsx"):
    print(f"Drafting Excel spreadsheet: {filename}...")
    
    # 1. Initialize Workbook
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    # Grab the sheets data from the JSON
    sheets_data = data.get("sheets", [])
    
    # Fallback in case the AI hallucinates the old single-sheet structure
    if not sheets_data and "headers" in data:
        sheets_data = [{"sheet_name": "Sheet1", "headers": data.get("headers", []), "rows": data.get("rows", [])}]

    # 2. Styling Definitions
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid") # Professional Blue
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # 3. Build Each Sheet
    for idx, sheet_data in enumerate(sheets_data):
        # Excel sheet names have a strict 31-character limit!
        sheet_name = sheet_data.get("sheet_name", f"Sheet {idx+1}")[:31]
        
        if idx == 0:
            ws = default_sheet
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(title=sheet_name)
            
        # Write Headers
        headers = sheet_data.get("headers", [])
        ws.append(headers)
        
        # Apply Header Styling & Freeze the top row
        for col_num, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        
        ws.freeze_panes = 'A2' 
        
        # Write Rows with "Zebra Striping" (Alternating Colors)
        rows = sheet_data.get("rows", [])
        for row_idx, row in enumerate(rows, start=2):
            ws.append(row)
            
            fill_color = "F2F2F2" if row_idx % 2 == 0 else "FFFFFF" # Light grey / White
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            
            for cell in ws[row_idx]:
                cell.fill = row_fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # 4. Auto-Adjust Column Widths based on content size
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            # Add a little padding to the longest word
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[col_letter].width = adjusted_width

    # 5. Save the masterpiece
    wb.save(filename)
    print("Excel spreadsheet complete!")